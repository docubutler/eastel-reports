"""Local inspection of which sheet holds the A2P (Q015) anchor and its layout."""
import sys
from openpyxl import load_workbook

paths = [
    r"2. CDR-Reconcialation-Report_template.xlsx",
    r"2-Invoice-Recon-postgres-report-Aug-26.xlsx",
]

for path in paths:
    print("\n=====================")
    print("FILE:", path)
    try:
        wb = load_workbook(path, data_only=False)
    except Exception as exc:
        print("  (could not open:", exc, ")")
        continue
    print("Sheets:", wb.sheetnames)
    for ws in wb.worksheets:
        hits = []
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and ("Q015" in cell.value or "A2P" in cell.value.lower()):
                    hits.append((cell.coordinate, repr(cell.value)))
        if hits:
            print(f"  -- sheet '{ws.title}' dims={ws.dimensions}")
            for coord, val in hits:
                print(f"       {coord}: {val}")
